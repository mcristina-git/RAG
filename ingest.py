import os
import shutil
import docx
from langchain_chroma import Chroma
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_core.documents import Document
from langchain_ollama import OllamaEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CARTELLA_DOCUMENTI = os.environ.get("RAG_DOCS_DIR", os.path.join(BASE_DIR, "Documenti"))
DB_PATH = os.environ.get("RAG_DB_PATH", os.path.join(BASE_DIR, "chroma_db"))
EMBED_MODEL = os.environ.get("OLLAMA_EMBED_MODEL", "nomic-embed-text")


def carica_docx(percorso):
    doc = docx.Document(percorso)
    docs_estratti = []
    nome_doc = os.path.basename(percorso)

    # Paragrafi normali (uniti e splittati normalmente)
    paragrafi = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    if paragrafi:
        testo_p = "\n".join(paragrafi)
        splitter = RecursiveCharacterTextSplitter(chunk_size=700, chunk_overlap=100)
        for chunk in splitter.split_text(testo_p):
            docs_estratti.append(Document(
                page_content=f"Documento: {nome_doc}\n{chunk}",
                metadata={"source": percorso}
            ))

    # Tabelle: ogni riga diventa un documento dedicato (così nessun requisito viene troncato)
    for t_idx, tabella in enumerate(doc.tables):
        for r_idx, riga in enumerate(tabella.rows):
            celle = [c.text.strip().replace("\n", " ") for c in riga.cells if c.text.strip()]
            celle_univoche = []
            for c in celle:
                if not celle_univoche or celle_univoche[-1] != c:
                    celle_univoche.append(c)
            if celle_univoche:
                riga_testo = " | ".join(celle_univoche)
                docs_estratti.append(Document(
                    page_content=f"Documento: {nome_doc} [Tabella]\n{riga_testo}",
                    metadata={"source": percorso}
                ))

    return docs_estratti


def indicizza_documenti():
    print("--- 1. Caricamento documenti... ---")
    chunks = []

    if not os.path.exists(CARTELLA_DOCUMENTI):
        print(f"Errore: la cartella {CARTELLA_DOCUMENTI} non esiste.")
        return

    for file in os.listdir(CARTELLA_DOCUMENTI):
        if file.startswith("~$") or file.startswith("."):
            continue

        percorso = os.path.join(CARTELLA_DOCUMENTI, file)
        estensione = file.lower().split(".")[-1]
        try:
            if estensione == "docx":
                chunks.extend(carica_docx(percorso))
            elif estensione == "pdf":
                pdf_docs = PyPDFLoader(percorso).load()
                splitter = RecursiveCharacterTextSplitter(chunk_size=700, chunk_overlap=100)
                chunks.extend(splitter.split_documents(pdf_docs))
            elif estensione == "txt":
                txt_docs = TextLoader(percorso, encoding="utf-8").load()
                splitter = RecursiveCharacterTextSplitter(chunk_size=700, chunk_overlap=100)
                chunks.extend(splitter.split_documents(txt_docs))
            print(f"[OK] Caricato: {file}")
        except Exception as e:
            print(f"[ERRORE] {file}: {e}")

    if not chunks:
        print(f"Nessun documento trovato in {CARTELLA_DOCUMENTI}.")
        return

    print(f"\nTotale frammenti creati: {len(chunks)}")

    print("\n--- 2. Salvataggio ChromaDB... ---")
    if os.path.exists(DB_PATH):
        shutil.rmtree(DB_PATH)

    embeddings = OllamaEmbeddings(model=EMBED_MODEL)
    batch_size = 50
    vectorstore = Chroma.from_documents(
        documents=chunks[:batch_size],
        embedding=embeddings,
        persist_directory=DB_PATH
    )

    for i in range(batch_size, len(chunks), batch_size):
        vectorstore.add_documents(chunks[i:i + batch_size])
        print(f"Indicizzati {min(i + batch_size, len(chunks))}/{len(chunks)}...")

    print("\n✅ Database rigenerato con successo!")


if __name__ == "__main__":
    indicizza_documenti()
